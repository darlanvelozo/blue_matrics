"""Testes dos exporters (Excel + HTML printable)."""
from __future__ import annotations

import io

from openpyxl import load_workbook

from apps.reports import exporters


def _kpi(current=100.0, previous=80.0, change_pct=25.0):
    return {"current": current, "previous": previous, "change_pct": change_pct}


SAMPLE_EXECUTIVE = {
    "period": {"start": "2025-05-01", "end": "2026-04-30"},
    "revenue": _kpi(100000.0, 80000.0, 25.0),
    "net_profit": _kpi(20000.0, 16000.0, 25.0),
    "avg_ticket": _kpi(2000.0, 2000.0, 0.0),
    "num_sales": _kpi(50, 40, 25.0),
    "overdue_rate": 2.5,
    "revenue_by_month": [
        {"month": "2026-03", "revenue": 8000.0, "sales_count": 4},
        {"month": "2026-04", "revenue": 9000.0, "sales_count": 5},
    ],
    "top_customers": [
        {"name": "Alpha LTDA", "total": 50000.0, "sales": 10},
        {"name": "Beta LTDA", "total": 30000.0, "sales": 5},
    ],
    "top_products": [
        {"name": "Pão", "total": 20000.0, "quantity": 1000.0},
    ],
}

SAMPLE_FINANCIAL = {
    "period": {"start": "2025-05-01", "end": "2026-04-30"},
    "cash_in": _kpi(50000.0, 40000.0, 25.0),
    "cash_out": _kpi(30000.0, 28000.0, 7.14),
    "net_profit": _kpi(20000.0, 12000.0, 66.67),
    "overdue_rate": 3.2,
    "cashflow_by_month": [
        {"month": "2026-03", "in": 5000.0, "out": 3000.0, "net": 2000.0},
        {"month": "2026-04", "in": 6000.0, "out": 3500.0, "net": 2500.0},
    ],
    "dre_monthly": [],
}

SAMPLE_COMMERCIAL = {
    "period": {"start": "2025-05-01", "end": "2026-04-30"},
    "revenue": _kpi(100000.0, 80000.0, 25.0),
    "num_sales": _kpi(50, 40, 25.0),
    "avg_ticket": _kpi(2000.0, 2000.0, 0.0),
    "revenue_by_month": [
        {"month": "2026-04", "revenue": 9000.0, "sales_count": 5},
    ],
    "top_customers": [
        {"name": "Alpha LTDA", "total": 50000.0, "sales": 10},
    ],
    "top_products": [
        {"name": "Pão", "total": 20000.0, "quantity": 1000.0},
    ],
    "by_salesperson": [
        {"name": "Maria", "total": 70000.0, "sales": 35},
        {"name": "João", "total": 30000.0, "sales": 15},
    ],
}


class TestExcelExporters:
    def test_executive_to_excel_is_xlsx(self):
        content = exporters.executive_to_excel(SAMPLE_EXECUTIVE, "Empresa Teste")
        assert isinstance(content, (bytes, bytearray))
        assert content[:2] == b"PK"
        wb = load_workbook(io.BytesIO(content))
        assert "Resumo" in wb.sheetnames
        assert "Top clientes" in wb.sheetnames

    def test_financial_to_excel(self):
        content = exporters.financial_to_excel(SAMPLE_FINANCIAL, "Empresa Teste")
        assert content[:2] == b"PK"
        wb = load_workbook(io.BytesIO(content))
        assert "Fluxo de caixa" in wb.sheetnames

    def test_commercial_to_excel(self):
        content = exporters.commercial_to_excel(SAMPLE_COMMERCIAL, "Empresa Teste")
        assert content[:2] == b"PK"
        wb = load_workbook(io.BytesIO(content))
        assert "Por vendedor" in wb.sheetnames


class TestHtmlExporters:
    def test_executive_to_html_renders(self):
        html = exporters.executive_to_html(SAMPLE_EXECUTIVE, "Empresa Teste")
        assert "<!doctype html>" in html.lower()
        assert "Empresa Teste" in html
        assert "Faturamento" in html
        assert "Alpha LTDA" in html
        # botão imprimir presente
        assert "window.print()" in html

    def test_financial_to_html_renders(self):
        html = exporters.financial_to_html(SAMPLE_FINANCIAL, "Empresa Teste")
        assert "Empresa Teste" in html
        assert "Fluxo de caixa" in html
        assert "Inadimplência" in html

    def test_commercial_to_html_renders(self):
        html = exporters.commercial_to_html(SAMPLE_COMMERCIAL, "Empresa Teste")
        assert "Maria" in html
        assert "Vendas por vendedor" in html
